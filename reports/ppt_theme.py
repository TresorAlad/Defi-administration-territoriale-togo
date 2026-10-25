"""Mise en page PowerPoint jury (charte Togo, 16:9)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

TOGO_GREEN = RGBColor(0, 106, 78)
TOGO_YELLOW = RGBColor(255, 206, 0)
TOGO_RED = RGBColor(210, 16, 52)
TOGO_INK = RGBColor(31, 41, 55)
TOGO_MUTED = RGBColor(107, 114, 128)
TOGO_BORDER = RGBColor(229, 231, 235)
TOGO_WHITE = RGBColor(255, 255, 255)
TOGO_PANEL = RGBColor(249, 250, 251)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_tricolor_band(slide, top=Inches(0), height=Inches(0.14)) -> None:
    w = SLIDE_W
    third = w / 3
    for idx, color in enumerate((TOGO_GREEN, TOGO_YELLOW, TOGO_RED)):
        shape = slide.shapes.add_shape(1, idx * third, top, third, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()


def add_footer(slide, text: str, slide_no: int | None = None) -> None:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(12), Inches(0.35))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(9)
    p.font.color.rgb = TOGO_MUTED
    if slide_no is not None:
        p2 = tf.add_paragraph()
        p2.text = str(slide_no)
        p2.alignment = PP_ALIGN.RIGHT
        p2.font.size = Pt(9)
        p2.font.color.rgb = TOGO_MUTED


def _textbox(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    size: int = 14,
    color=TOGO_INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = align
    p.line_spacing = 1.15


def slide_cover(prs: Presentation, title: str, subtitle: str, tagline: str) -> None:
    slide = _blank_slide(prs)
    add_tricolor_band(slide, height=Inches(0.18))
    panel = slide.shapes.add_shape(1, Inches(0.55), Inches(1.05), Inches(12.2), Inches(5.35))
    panel.fill.solid()
    panel.fill.fore_color.rgb = TOGO_PANEL
    panel.line.color.rgb = TOGO_BORDER

    _textbox(
        slide,
        Inches(0.9),
        Inches(1.55),
        Inches(11.5),
        Inches(1.4),
        title,
        size=34,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    _textbox(
        slide,
        Inches(1.1),
        Inches(2.95),
        Inches(11.1),
        Inches(0.9),
        subtitle,
        size=16,
        color=TOGO_MUTED,
        align=PP_ALIGN.CENTER,
    )
    _textbox(
        slide,
        Inches(1.2),
        Inches(4.05),
        Inches(10.9),
        Inches(1.2),
        tagline,
        size=13,
        color=TOGO_INK,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, "Data Challenge Togo · Pipeline Bronze → Gold · Dashboard Streamlit", 1)


def slide_kpi_context(
    prs: Presentation,
    intro: str,
    kpis: list[tuple[str, str]],
    insight: str,
    slide_no: int,
) -> None:
    slide = _blank_slide(prs)
    add_tricolor_band(slide)
    _textbox(slide, Inches(0.55), Inches(0.28), Inches(12), Inches(0.55), "Contexte et enjeux", size=26, bold=True)

    x0 = Inches(0.55)
    w = Inches(2.2)
    gap = Inches(0.12)
    for i, (label, value) in enumerate(kpis):
        left = x0 + i * (w + gap)
        card = slide.shapes.add_shape(1, left, Inches(1.05), w, Inches(1.05))
        card.fill.solid()
        card.fill.fore_color.rgb = TOGO_WHITE
        card.line.color.rgb = TOGO_BORDER
        accent = slide.shapes.add_shape(1, left, Inches(1.05), Inches(0.07), Inches(1.05))
        accent.fill.solid()
        accent.fill.fore_color.rgb = (TOGO_GREEN, TOGO_YELLOW, TOGO_RED)[i % 3]
        accent.line.fill.background()
        _textbox(slide, left + Inches(0.15), Inches(1.15), w - Inches(0.2), Inches(0.35), value, size=20, bold=True)
        _textbox(slide, left + Inches(0.15), Inches(1.65), w - Inches(0.2), Inches(0.45), label, size=10, color=TOGO_MUTED)

    _textbox(slide, Inches(0.55), Inches(2.35), Inches(12.2), Inches(1.35), intro, size=13)
    box = slide.shapes.add_shape(1, Inches(0.55), Inches(3.85), Inches(12.2), Inches(2.85))
    box.fill.solid()
    box.fill.fore_color.rgb = TOGO_WHITE
    box.line.color.rgb = TOGO_GREEN
    _textbox(slide, Inches(0.75), Inches(4.05), Inches(11.8), Inches(2.5), insight, size=14, bold=True)
    add_footer(slide, "Profils territoriaux · Inégalités d'accès aux services publics", slide_no)


def slide_pipeline(prs: Presentation, slide_no: int) -> None:
    slide = _blank_slide(prs)
    add_tricolor_band(slide)
    _textbox(slide, Inches(0.55), Inches(0.28), Inches(12), Inches(0.55), "Méthodologie et fiabilité", size=26, bold=True)
    steps = [
        ("Bronze", "Exports open data bruts\n(CSV recensement, CEET, TdE, état civil)"),
        ("Silver", "Nettoyage, typage, flags\n(homonymes, séries partielles)"),
        ("Ref", "39 préfectures, 117 communes\nAlias libellés (Tandjoaré, Agouè-Nyivé…)"),
        ("Gold", "Marts analytiques, scores,\nrecommandations 39/39"),
    ]
    x = Inches(0.45)
    w = Inches(2.85)
    for i, (title, body) in enumerate(steps):
        left = x + i * (w + Inches(0.25))
        card = slide.shapes.add_shape(1, left, Inches(1.15), w, Inches(2.35))
        card.fill.solid()
        card.fill.fore_color.rgb = TOGO_WHITE
        card.line.color.rgb = TOGO_BORDER
        _textbox(slide, left + Inches(0.15), Inches(1.3), w - Inches(0.25), Inches(0.4), title, size=16, bold=True, color=TOGO_GREEN)
        _textbox(slide, left + Inches(0.15), Inches(1.85), w - Inches(0.25), Inches(1.45), body, size=12, color=TOGO_INK)
        if i < len(steps) - 1:
            arrow = slide.shapes.add_textbox(left + w + Inches(0.02), Inches(2.05), Inches(0.25), Inches(0.4))
            arrow.text_frame.paragraphs[0].text = "→"
            arrow.text_frame.paragraphs[0].font.size = Pt(22)
            arrow.text_frame.paragraphs[0].font.color.rgb = TOGO_MUTED

    _textbox(
        slide,
        Inches(0.55),
        Inches(3.85),
        Inches(12.2),
        Inches(2.85),
        "Reproductibilité : python3 data/bronze/clean_data.py puis python3 -m pipeline.gold.build. "
        "Limites documentées : découpage administratif 2010 vs 2022, anomalies source (>100 % électricité sur certains territoires).",
        size=13,
    )
    add_footer(slide, "Jointure territoriale ref + alias · Traces silver", slide_no)


def slide_chart(
    prs: Presentation,
    section: str,
    title: str,
    narrative: str,
    image_path: Path,
    slide_no: int,
    *,
    image_left=Inches(0.55),
    image_top=Inches(2.05),
    image_width=Inches(12.2),
    image_height=Inches(4.85),
) -> None:
    slide = _blank_slide(prs)
    add_tricolor_band(slide)
    _textbox(slide, Inches(0.55), Inches(0.22), Inches(12), Inches(0.35), section, size=11, bold=True, color=TOGO_GREEN)
    _textbox(slide, Inches(0.55), Inches(0.48), Inches(12), Inches(0.55), title, size=24, bold=True)
    _textbox(slide, Inches(0.55), Inches(1.12), Inches(12.2), Inches(0.85), narrative, size=12, color=TOGO_MUTED)
    slide.shapes.add_picture(str(image_path), image_left, image_top, width=image_width, height=image_height)
    add_footer(slide, "Visualisation Plotly · Même charte que le dashboard Streamlit", slide_no)


def slide_two_charts(
    prs: Presentation,
    section: str,
    title: str,
    narrative: str,
    left_img: Path,
    right_img: Path,
    slide_no: int,
) -> None:
    slide = _blank_slide(prs)
    add_tricolor_band(slide)
    _textbox(slide, Inches(0.55), Inches(0.22), Inches(12), Inches(0.35), section, size=11, bold=True, color=TOGO_GREEN)
    _textbox(slide, Inches(0.55), Inches(0.48), Inches(12), Inches(0.55), title, size=24, bold=True)
    _textbox(slide, Inches(0.55), Inches(1.05), Inches(12.2), Inches(0.7), narrative, size=12, color=TOGO_MUTED)
    slide.shapes.add_picture(str(left_img), Inches(0.45), Inches(1.85), width=Inches(6.35), height=Inches(4.95))
    slide.shapes.add_picture(str(right_img), Inches(6.95), Inches(1.85), width=Inches(6.35), height=Inches(4.95))
    add_footer(slide, "État civil et identité · Sources silver/gold", slide_no)


def slide_chart_and_map(
    prs: Presentation,
    section: str,
    title: str,
    narrative: str,
    chart_path: Path,
    map_path: Path | None,
    slide_no: int,
) -> None:
    slide = _blank_slide(prs)
    add_tricolor_band(slide)
    _textbox(slide, Inches(0.55), Inches(0.22), Inches(12), Inches(0.35), section, size=11, bold=True, color=TOGO_GREEN)
    _textbox(slide, Inches(0.55), Inches(0.48), Inches(12), Inches(0.55), title, size=24, bold=True)
    _textbox(slide, Inches(0.55), Inches(1.05), Inches(12.2), Inches(0.65), narrative, size=12, color=TOGO_MUTED)
    slide.shapes.add_picture(str(chart_path), Inches(0.45), Inches(1.78), width=Inches(7.35), height=Inches(5.05))
    if map_path and map_path.exists():
        slide.shapes.add_picture(str(map_path), Inches(8.0), Inches(1.78), width=Inches(5.0), height=Inches(5.05))
    add_footer(slide, "Services publics · Électricité et agences CEET / TdE", slide_no)


def slide_recommandations(
    prs: Presentation,
    chart_path: Path,
    fiches: list[tuple[str, float, str]],
    slide_no: int,
) -> None:
    slide = _blank_slide(prs)
    add_tricolor_band(slide)
    _textbox(slide, Inches(0.55), Inches(0.22), Inches(12), Inches(0.35), "Pilotage territorial", size=11, bold=True, color=TOGO_GREEN)
    _textbox(slide, Inches(0.55), Inches(0.48), Inches(12), Inches(0.55), "Recommandations prioritaires", size=24, bold=True)
    slide.shapes.add_picture(str(chart_path), Inches(0.45), Inches(1.15), width=Inches(6.5), height=Inches(5.65))

    y = Inches(1.25)
    for name, score, reco in fiches[:3]:
        card = slide.shapes.add_shape(1, Inches(7.15), y, Inches(5.75), Inches(1.65))
        card.fill.solid()
        card.fill.fore_color.rgb = TOGO_WHITE
        card.line.color.rgb = TOGO_BORDER
        _textbox(
            slide,
            Inches(7.35),
            y + Inches(0.12),
            Inches(5.4),
            Inches(0.35),
            f"{name} · score {score:.1f}",
            size=13,
            bold=True,
            color=TOGO_GREEN,
        )
        short = reco if len(reco) <= 220 else reco[:217] + "..."
        _textbox(slide, Inches(7.35), y + Inches(0.48), Inches(5.4), Inches(1.05), short, size=10, color=TOGO_INK)
        y += Inches(1.78)

    add_footer(slide, "mart_recommandations.csv · Score composite gold", slide_no)


def slide_conclusion(prs: Presentation, bullets: list[str], slide_no: int) -> None:
    slide = _blank_slide(prs)
    add_tricolor_band(slide)
    _textbox(slide, Inches(0.55), Inches(0.28), Inches(12), Inches(0.55), "Conclusion et livrables", size=26, bold=True)
    y = Inches(1.15)
    for line in bullets:
        dot = slide.shapes.add_shape(1, Inches(0.65), y + Inches(0.08), Inches(0.12), Inches(0.12))
        dot.fill.solid()
        dot.fill.fore_color.rgb = TOGO_GREEN
        dot.line.fill.background()
        _textbox(slide, Inches(0.95), y, Inches(11.8), Inches(0.75), line, size=15)
        y += Inches(0.82)

    _textbox(
        slide,
        Inches(0.55),
        Inches(5.85),
        Inches(12.2),
        Inches(0.55),
        "GitHub : TresorAlad/Defi-administration-territoriale-togo · Dashboard Streamlit déployé",
        size=12,
        color=TOGO_MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, "Merci · Questions du jury", slide_no)
