#!/usr/bin/env python3
"""Genere le PowerPoint jury (10 slides max) depuis data/gold/."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt

PROJECT_ROOT = Path(__file__).resolve().parents[1]
GOLD = PROJECT_ROOT / "data" / "gold"
OUT = PROJECT_ROOT / "reports" / "rapport_jury.pptx"


def add_title(slide, title: str, subtitle: str = "") -> None:
    slide.shapes.title.text = title
    if subtitle and len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle


def add_bullets(slide, lines: list[str]) -> None:
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, line in enumerate(lines):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(18)


def main() -> None:
    synthese = pd.read_csv(GOLD / "mart_synthese_nationale.csv")
    ineg = pd.read_csv(GOLD / "mart_inegalites.csv")
    reco = pd.read_csv(GOLD / "mart_recommandations.csv")
    evol = pd.read_csv(GOLD / "mart_evolution_2010_2022.csv")

    kpi = dict(zip(synthese["indicateur"], synthese["valeur"]))
    ineg_d = dict(zip(ineg["indicateur"], ineg["valeur"]))

    prs = Presentation()
    # 1
    s = prs.slides.add_slide(prs.slide_layouts[0])
    add_title(
        s,
        "Territoires et inégalités d'accès au Togo",
        "Data Challenge - Bronze / Silver / Ref / Gold - Dashboard Python",
    )
    # 2
    s = prs.slides.add_slide(prs.slide_layouts[1])
    add_title(s, "Contexte et enjeux")
    add_bullets(
        s,
        [
            "39 préfectures, 117 communes, poids démographique très inégal.",
            f"Grand Lomé (Golfe + Agouè-Nyivé) : {ineg_d.get('part_grand_lome_golfe_agoe_pct', 0):.1f} % de la population.",
            f"Écart max/min entre communes : {ineg_d.get('ratio_pop_max_min_commune', 0):.0f}x.",
            "Objectif : rapprocher l'administration des citoyens via la donnée.",
        ],
    )
    # 3
    s = prs.slides.add_slide(prs.slide_layouts[1])
    add_title(s, "Méthodologie")
    add_bullets(
        s,
        [
            "Bronze : open data brutes | Silver : nettoyage | Ref : référentiel admin + alias.",
            "Gold : marts analytiques (39 préfectures, scores, recommandations).",
            "Pipeline reproductible : clean_data.py + pipeline.gold.build.",
            "Limites documentées : homonymes RGPH, découpage 2010/2022.",
        ],
    )
    # 4
    s = prs.slides.add_slide(prs.slide_layouts[1])
    add_title(s, "Démographie 2022")
    add_bullets(
        s,
        [
            f"Population nationale : {kpi.get('population_nationale_2022', 0):,.0f} habitants.",
            "Profils par préfecture et commune (poids national, rang).",
            "Visualisation interactive : pyramide territoriale, filtres région.",
        ],
    )
    # 5
    n_comp = len(evol[evol["comparable_decoupage"] == "oui"])
    top_evol = evol[evol["comparable_decoupage"] == "oui"].nlargest(3, "variation_pct")
    evol_lines = [f"{n_comp}/39 préfectures comparables 2010-2022 (libellés alignés)."]
    for _, r in top_evol.iterrows():
        evol_lines.append(f"{r['prefecture_nom']} : +{r['variation_pct']} %")
    s = prs.slides.add_slide(prs.slide_layouts[1])
    add_title(s, "Évolution 2010-2022")
    add_bullets(s, evol_lines)
    # 6
    s = prs.slides.add_slide(prs.slide_layouts[1])
    add_title(s, "Accès à l'électricité")
    add_bullets(
        s,
        [
            f"Taux min préfecture : {ineg_d.get('taux_electricite_min_prefecture', 0):.1f} %.",
            f"Taux max préfecture : {ineg_d.get('taux_electricite_max_prefecture', 0):.1f} %.",
            "Écarts régionaux majeurs : Lomé vs Savanes / Kara.",
        ],
    )
    # 7
    s = prs.slides.add_slide(prs.slide_layouts[1])
    add_title(s, "Agences CEET et TdE")
    add_bullets(
        s,
        [
            f"{int(kpi.get('nb_agences_ceet', 0))} agences CEET, {int(kpi.get('nb_agences_tde', 0))} TdE.",
            "Ratio habitants/agence : indicateur de pression d'accès.",
            "Cartographie WKT des points de service.",
        ],
    )
    # 8
    s = prs.slides.add_slide(prs.slide_layouts[1])
    add_title(s, "État civil et identité")
    add_bullets(
        s,
        [
            "Suivi des naissances enregistrées (< 45 jours).",
            "Production CNI, passeports, cartes professionnelles.",
            "Série 2021 partielle signalée en silver.",
        ],
    )
    # 9
    top3 = reco.head(3)
    reco_lines = ["Score composite : électricité + agences + croissance."]
    for _, r in top3.iterrows():
        reco_lines.append(f"{r['prefecture_nom']} (score {r['score_priorite']})")
    s = prs.slides.add_slide(prs.slide_layouts[1])
    add_title(s, "Recommandations prioritaires")
    add_bullets(s, reco_lines)
    # 10
    s = prs.slides.add_slide(prs.slide_layouts[1])
    add_title(s, "Conclusion")
    add_bullets(
        s,
        [
            "Gold 39/39 : jointures ref + alias open data.",
            "Dashboard Streamlit + rapport PPT alignés sur les marts.",
            "Prochaine étape : déploiement cloud + mise à jour trimestrielle.",
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Rapport genere : {OUT}")


if __name__ == "__main__":
    main()
